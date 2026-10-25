import io
import os
import sys
import threading

import pikepdf
import pymupdf
import pytest
from conftest import PASSWORD, make_pdf

from pdfmanager.core import compare, convert, ocr, optimize, security, text
from pdfmanager.core.document import Cancelled, PdfDocument, PdfError
from pdfmanager.services import external, jobs


class FakeTask:
    def __init__(self):
        self.cancelled = False
        self.progress = []

    def emit_partial(self, value):
        self.progress.append(value)


@pytest.fixture
def doc(sample_pdf):
    return PdfDocument(sample_pdf)


def scanned_pdf(path, rotation=0):
    """Page « scannée » : du texte devenu image, droite à l'écran (page éventuellement
    tournée et recadrée, comme un vrai scan)."""
    src = pymupdf.open()
    p = src.new_page()
    p.insert_text((72, 100), "Facture numéro 2026-041", fontsize=16)
    p.insert_text((72, 140), "Montant total : 1 250,00 € échéance été", fontsize=12)
    img = p.get_pixmap(dpi=200).tobytes("png")
    d = pymupdf.open()
    page = d.new_page()
    page.set_cropbox(pymupdf.Rect(5, 10, 590, 832))
    page.set_rotation(rotation)
    page.insert_image(page.rect * page.derotation_matrix, stream=img, rotate=rotation)
    d.new_page().insert_text((72, 72), "Page avec du vrai texte")
    d.save(path)
    return path


# ----- sécurité et propriétés ------------------------------------------------------------

def readable(doc):
    return "Bonjour page 1" in " ".join(w[4] for w in text.page_words(doc, 0)[:3])


def test_protect_open_password_and_undo(doc, tmp_path):
    security.protect(doc, "ouvrir", "proprio", {"print"})
    assert security.is_protected(doc) and readable(doc)
    doc.save(tmp_path / "p.pdf")
    again = PdfDocument(tmp_path / "p.pdf")
    assert again.needs_password and not again.unlock("faux") and again.unlock("ouvrir")
    assert security.allowed(again) == {"print"} and readable(again)
    with pytest.raises(PdfError, match="permissions"):
        security.unprotect(again)
    # Régression : un mot de passe faux (ou needs_pass) sur un document déverrouillé
    # remplaçait la clé de déchiffrement de PyMuPDF : contenu illisible.
    assert not security.has_owner_rights(again, "faux")
    assert again.unlock("faux") and again.encrypted and readable(again)
    assert security.has_owner_rights(again, "proprio")
    security.unprotect(again)
    again.save(tmp_path / "libre.pdf")
    libre = PdfDocument(tmp_path / "libre.pdf")
    assert not libre.needs_password and readable(libre)
    doc.undo()
    assert not security.is_protected(doc) and readable(doc)


def test_restrictions_only_and_validation(doc, tmp_path):
    with pytest.raises(ValueError, match="différer"):
        security.protect(doc, "x", "x", set())
    with pytest.raises(ValueError, match="mot de passe des permissions"):
        security.protect(doc, "", "", {"print"})
    security.protect(doc, "", "proprio", {"print", "copy"})
    doc.save(tmp_path / "r.pdf")
    again = PdfDocument(tmp_path / "r.pdf")
    assert not again.needs_password and security.allowed(again) == {"print", "copy"}


def test_encrypted_document_keeps_password_through_edits(encrypted_pdf, tmp_path):
    doc = PdfDocument(encrypted_pdf)
    doc.unlock(PASSWORD)
    security.set_metadata(doc, {"title": "Été", "author": "Moi"})
    doc.save(tmp_path / "m.pdf")
    again = PdfDocument(tmp_path / "m.pdf")
    assert again.unlock(PASSWORD) and security.metadata(again)["title"] == "Été"
    doc.undo()
    assert security.metadata(doc)["title"] == ""


# ----- OCR ---------------------------------------------------------------------------------

@pytest.mark.parametrize("rotation", [0, 90, 270])
def test_ocr_adds_searchable_text_where_it_is_seen(tmp_path, rotation):
    doc = PdfDocument(scanned_pdf(tmp_path / "scan é.pdf", rotation))
    assert ocr.pages_to_process(doc, [0, 1], all_pages=False) == [0]
    before = doc.render(0, 1).samples
    task = FakeTask()
    assert jobs.ocr_document(task, doc, [0], "fra+eng") == 1
    assert task.progress[0][:2] == (1, 1)
    words = {w[4] for w in text.page_words(doc, 0)}
    assert {"Facture", "numéro", "échéance"} <= words
    # le mot invisible recouvre le mot de l'image (repère de l'écran)
    page = doc.doc[0]
    found = next(w for w in text.page_words(doc, 0) if w[4] == "Facture")
    on_screen = pymupdf.Rect(found[:4]) * page.rotation_matrix
    pix = doc.render(0, 1)
    dark = [(x, y) for y in range(int(on_screen.y0), int(on_screen.y1))
            for x in range(int(on_screen.x0), int(on_screen.x1)) if pix.pixel(x, y)[0] < 100]
    assert len(dark) > 30
    assert pix.samples == before  # aspect inchangé
    doc.undo()
    assert text.page_words(doc, 0) == []


def test_ocr_cancel(tmp_path):
    doc = PdfDocument(scanned_pdf(tmp_path / "scan.pdf"))
    task = FakeTask()
    task.cancelled = True
    with pytest.raises(Cancelled):
        jobs.ocr_document(task, doc, [0], "fra")
    assert not doc.modified


# ----- compression, PDF/A, réparation --------------------------------------------------------

def photo_pdf(path):
    d = pymupdf.open()
    # bruit : incompressible sans perte, comme le grain d'une photo
    pix = pymupdf.Pixmap(pymupdf.csRGB, 800, 600, os.urandom(800 * 600 * 3), False)
    d.new_page().insert_image(pymupdf.Rect(50, 50, 350, 275), pixmap=pix)
    d.save(path)
    return path


def test_compress_levels(tmp_path):
    doc = PdfDocument(photo_pdf(tmp_path / "photo.pdf"))
    before = len(doc.current_bytes())
    lossless = optimize.compress(doc.current_bytes(), None, "lossless", linearize=True)
    strong = jobs.compress(FakeTask(), doc, "strong", linearize=False)
    assert len(strong) < before / 3 and len(strong) < len(lossless)
    with pikepdf.open(io.BytesIO(lossless)) as pdf:
        assert pdf.is_linearized
    assert pymupdf.open("pdf", strong)[0].get_images()  # l'image est toujours là


@pytest.mark.skipif(external.find_ghostscript() is None, reason="Ghostscript absent")
def test_pdfa_with_ghostscript(doc):
    data, _notes = external.to_pdfa(FakeTask(), doc.current_bytes(),
                                    external.find_ghostscript())
    out = pymupdf.open("pdf", data)
    assert "pdfaid:part" in out.get_xml_metadata()
    assert "Bonjour" in out[0].get_text()


def test_pdfa_refused_rather_than_faked(doc):
    with pytest.raises(PdfError, match="Ghostscript"):  # police non intégrée
        external.to_pdfa(FakeTask(), doc.current_bytes(), None)


def test_repair_broken_xref(tmp_path):
    good = make_pdf(tmp_path / "ok.pdf").read_bytes()
    broken = tmp_path / "abîmé.pdf"
    broken.write_bytes(good[:good.rfind(b"xref")] + b"%%EOF garbage")
    doc = PdfDocument(broken)
    assert doc.repaired and doc.page_count == 3 and not doc.incremental


def test_repair_falls_back_to_qpdf(tmp_path, monkeypatch):
    path = make_pdf(tmp_path / "x.pdf")
    real_open = pymupdf.open
    calls = []

    def flaky(*a, **k):  # MuPDF refuse le fichier d'origine, accepte la version de qpdf
        calls.append(1)
        if len(calls) == 1:
            raise pymupdf.FileDataError("cannot open broken document")
        return real_open(*a, **k)

    monkeypatch.setattr(pymupdf, "open", flaky)
    doc = PdfDocument(path)
    assert doc.repaired and doc.page_count == 3
    (tmp_path / "rien.pdf").write_bytes(b"pas un pdf")
    with pytest.raises(PdfError, match="pas un PDF valide"):
        PdfDocument(tmp_path / "rien.pdf")


# ----- conversions -------------------------------------------------------------------------

def test_to_images_text_and_back(doc, tmp_path):
    targets = convert.image_targets(doc, [0, 2], tmp_path, "jpg")
    assert targets[1].name == "Été déjà vu - page 003.jpg"
    convert.to_images(doc, [0, 2], targets, 72)
    assert all(t.exists() for t in targets)
    pdf = pymupdf.open("pdf", convert.images_to_pdf([t.read_bytes() for t in targets]))
    assert pdf.page_count == 2 and pdf[0].get_images()
    txt = convert.to_text(doc)
    assert "Bonjour page 1" in txt and txt.count("\f") == 2


def test_to_docx(encrypted_pdf, tmp_path):
    from docx import Document

    doc = PdfDocument(encrypted_pdf)  # document protégé : cas qui échouait
    doc.unlock(PASSWORD)
    target = tmp_path / "Été.docx"
    jobs.to_docx(FakeTask(), doc, target)
    paragraphs = " ".join(p.text for p in Document(str(target)).paragraphs)
    assert "Bonjour page 1" in paragraphs and "renard" in paragraphs


@pytest.mark.skipif(os.name == "nt", reason="faux LibreOffice en script shell")
def test_office_to_pdf_with_fake_libreoffice(tmp_path):
    """Arguments passés à LibreOffice : profil jetable, sortie dans le dossier de travail."""
    fake = tmp_path / "soffice"
    fake.write_text(f"""#!{sys.executable}
import sys, pymupdf
from pathlib import Path
args = sys.argv[1:]
assert args[0].startswith("-env:UserInstallation=file://") and "--headless" in args
outdir, src = Path(args[args.index("--outdir") + 1]), Path(args[-1])
d = pymupdf.open(); d.new_page().insert_text((72, 72), src.read_text())
d.save(outdir / (src.stem + ".pdf"))
""")
    fake.chmod(0o755)
    source = tmp_path / "Lettre été.txt"
    source.write_text("Bonjour")
    data = external.office_to_pdf(FakeTask(), source, fake)
    assert "Bonjour" in pymupdf.open("pdf", data)[0].get_text()


@pytest.mark.skipif(os.name == "nt", reason="commande POSIX")
def test_external_run_is_cancellable():
    task = FakeTask()
    threading.Timer(0.3, lambda: setattr(task, "cancelled", True)).start()
    with pytest.raises(Cancelled):
        external.run(task, ["sleep", "30"])
    with pytest.raises(PdfError, match="code 3"):
        external.run(FakeTask(), ["sh", "-c", "echo oups; exit 3"])


# ----- comparaison -------------------------------------------------------------------------

def test_compare_documents(tmp_path):
    a = PdfDocument(make_pdf(tmp_path / "a.pdf", pages=2))
    d = pymupdf.open(a.path)
    d[1].add_redact_annot(d[1].search_for("brun")[0])
    d[1].apply_redactions()
    d[1].insert_text((72, 200), "Ajout final", fontsize=12)
    d.save(tmp_path / "b.pdf")
    b = PdfDocument(tmp_path / "b.pdf")
    changes = compare.compare(compare.document_words(a), compare.document_words(b))
    kinds = {(c.kind, c.old_text, c.new_text) for c in changes}
    assert ("removed", "brun", "") in kinds and ("added", "", "Ajout final") in kinds
    assert all(page == 1 for c in changes for page, _, _ in c.old + c.new)
    assert compare.compare(compare.document_words(a), compare.document_words(a)) == []


def test_edit_text_recognized_by_ocr(tmp_path):
    """Régression : modifier un texte reconnu par OCR écrivait le nouveau texte par-dessus
    l'ancien, resté visible dans l'image. La ligne est maintenant effacée de l'image (couleur
    du papier) et réécrite en clair (couleur de l'encre)."""
    from pdfmanager.core import editing as E

    src = pymupdf.open()
    p = src.new_page()
    p.draw_rect(p.rect, color=None, fill=(0.97, 0.93, 0.85))  # papier crème
    p.insert_text((72, 100), "Facture numéro 2026-041", fontsize=16, color=(0.1, 0.2, 0.6))
    img = p.get_pixmap(dpi=200).tobytes("png")
    d = pymupdf.open()
    d.new_page().insert_image(d[0].rect, stream=img)
    d.save(tmp_path / "scan.pdf")
    doc = PdfDocument(tmp_path / "scan.pdf")
    jobs.ocr_document(FakeTask(), doc, [0], "fra")
    block = E.text_block_at(doc, 0, pymupdf.Point(100, 95))
    assert block.lines[0].invisible and not block.lines[0].mixed
    warnings = E.edit_block(doc, block, block.text.replace("2026-041", "2026-099"))
    assert any("OCR" in w for w in warnings)
    assert "2026-099" in [w[4] for w in text.page_words(doc, 0)]
    pix = doc.render(0, 2)
    ink = [pix.pixel(x, y) for y in range(170, 205) for x in range(140, 460)
           if sum(pix.pixel(x, y)) < 450]
    assert ink and all(b > r + 40 for r, _g, b in ink)  # encre bleue, rien de noir dessous
    assert sum(pix.pixel(150, 160)) > 650  # fond crème conservé (pas de bande blanche/noire)
    assert pix.pixel(150, 160) != (255, 255, 255)
    assert len(ink) < 3000  # une seule ligne de texte, pas deux superposées


# ----- traduction (DeepL remplacé par un traducteur factice) ----------------------------

def _upper(texts):
    return [t.upper() for t in texts]


def test_deepl_batches_and_errors(monkeypatch):
    from pdfmanager.core import deepl

    texts = [f"phrase {i}" for i in range(120)]
    assert deepl.batches(texts) == [(0, 50), (50, 100), (100, 120)]
    assert deepl.batches(["x" * 60_000, "y" * 60_000, "z"]) == [(0, 1), (1, 3)]
    assert deepl.endpoint("abc:fx").startswith("https://api-free.deepl.com")
    assert deepl.parse_response({"translations": [{"text": "Hello"}]}, 1) == ["Hello"]
    with pytest.raises(deepl.DeepLError):
        deepl.parse_response({"translations": []}, 1)
    with pytest.raises(deepl.DeepLError, match="clé"):
        deepl.translate("  ", ["Bonjour"], "EN-US")
    sent = []

    class Reply(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def fake_urlopen(request, timeout):
        import json
        body = json.loads(request.data)
        sent.append((request.full_url, request.get_header("Authorization"), body))
        reply = {"translations": [{"text": t.upper()} for t in body["text"]]}
        return Reply(json.dumps(reply).encode())

    monkeypatch.setattr(deepl, "urlopen", fake_urlopen)
    assert deepl.translate("k:fx", ["un", "deux"], "EN-US") == ["UN", "DEUX"]
    assert sent[0][0].startswith("https://api-free") and sent[0][1] == "DeepL-Auth-Key k:fx"
    assert sent[0][2] == {"text": ["un", "deux"], "target_lang": "EN-US"}


def test_translate_in_place(sample_pdf):
    from pdfmanager.core import translate

    data = open(sample_pdf, "rb").read()
    sent = []
    out, report = translate.translate_pdf(data, None, lambda t: sent.extend(t) or _upper(t))
    assert "Bonjour page 1" in sent and report.translated >= 6 and report.overflow == 0
    before, after = pymupdf.open(stream=data), pymupdf.open(stream=out)
    assert "BONJOUR PAGE 1" in after[0].get_text() and "Bonjour" not in after[0].get_text()
    w0, w1 = before[0].get_text("words")[0], after[0].get_text("words")[0]
    assert abs(w0[0] - w1[0]) < 1.5 and abs(w0[1] - w1[1]) < 3  # même place


def test_translate_table_and_pages(tmp_path):
    from pdfmanager.core import translate

    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 80), "Facture n° 2025-042", fontsize=16, fontname="hebo")
    for y, row in ((150, ("Désignation", "Qté", "Total")), (175, ("Conseil", "3", "1 350,00"))):
        for x, cell in zip((77, 277, 437), row, strict=True):
            p.insert_text((x, y), cell, fontsize=10)
    doc.new_page().insert_text((72, 80), "Deuxième page", fontsize=12)
    data = doc.tobytes()
    sent = []
    out, report = translate.translate_pdf(data, [0], lambda t: sent.extend(t) or _upper(t))
    assert all(translate.has_words(t) for t in sent), sent  # nombres non envoyés
    after = pymupdf.open(stream=out)
    words = {w[4]: w for w in after[0].get_text("words")}
    assert {"DÉSIGNATION", "QTÉ", "TOTAL", "CONSEIL", "1", "350,00"} <= set(words)
    assert abs(words["TOTAL"][0] - 437) < 1.5  # chaque cellule à sa place
    assert after[1].get_text().strip() == "Deuxième page"  # page non demandée intacte
    assert report.overflow == 0


def test_translate_pieces():
    from pdfmanager.core import translate

    assert translate.join("exem-", "ple") == "exemple"
    assert translate.join("exemple", "Suite") == "exemple Suite"
    assert translate.winansi("Été 10 € « oui »") and not translate.winansi("Привет")
    assert not translate.has_words("12/03/2026 1 250,00") and translate.has_words("Total HT")
